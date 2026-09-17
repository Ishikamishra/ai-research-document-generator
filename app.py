import streamlit as st
from docx import Document
from docx.shared import Pt
from datetime import datetime
import os
import io
from config import Config
from ai_provider import get_ai_provider
from news_search import search_recent_news as fetch_recent_news

# Page config
st.set_page_config(
    page_title="AI Research Document Generator",
    page_icon="📊",
    layout="wide"
)

# Display configuration info in sidebar (optional debug info)
with st.sidebar:
    st.subheader("🤖 AI Provider Selection")
    
    # Provider selector
    provider_options = {
        "Claude (claude-3-haiku)": "claude",
        "OpenAI (gpt-4o-mini)": "openai",
        "Gemini (gemini-3.6-flash)": "gemini"
    }
    
    selected_label = st.selectbox(
        "Choose AI Provider:",
        options=list(provider_options.keys()),
        index=list(provider_options.values()).index(Config.AI_PROVIDER) if Config.AI_PROVIDER in provider_options.values() else 2
    )
    selected_provider = provider_options[selected_label]
    
    # Show API key status
    st.markdown("*API Key Status:*")
    col1, col2, col3 = st.columns(3)
    with col1:
        status = "✅" if Config.CLAUDE_API_KEY else "❌"
        st.metric("Claude", status)
    with col2:
        status = "✅" if Config.OPENAI_API_KEY else "❌"
        st.metric("OpenAI", status)
    with col3:
        status = "✅" if Config.GEMINI_API_KEY else "❌"
        st.metric("Gemini", status)
    
    if st.checkbox("🔧 Show Configuration", value=False):
        st.subheader("Current Configuration")
        config_info = Config.display_config()
        for key, value in config_info.items():
            st.code(f"{key}: {value}")

# Title
st.title("📊 AI Research Document Generator")
st.markdown("*Generate professional company research documents in 2-3 minutes!*")
st.markdown("---")

# Sidebar
with st.sidebar:
    st.header("ℹ️ About")
    st.markdown("""
    This tool generates comprehensive research documents including:
    - Executive Summary
    - Company Overview
    - Business Segments & Brands
    - Sales & Distribution Model
    - GCC Market Analysis
    - AI & Digital Transformation
    - Technology Stack
    - Strategic Transformation
    - Recent News (2024-2026)
    - Key Insights & Challenges
    
    *Powered by:* Claude, OpenAI, or Gemini APIs + DuckDuckGo Search
    """)
    
    st.header("📝 Instructions")
    st.markdown("""
    1. Select AI provider
    2. Enter company name
    3. Click 'Generate Research'
    4. Wait 2-3 minutes
    5. Download the document
    """)

# Main content
company_name = st.text_input(
    "Enter Company Name:",
    placeholder="e.g., Microsoft, Diageo, Salesforce",
    help="Enter the name of the company you want to research"
)

def search_recent_news(company_name, max_results=None):
    """Search for recent news using DuckDuckGo"""
    if max_results is None:
        max_results = Config.MAX_NEWS_RESULTS
    
    try:
        return fetch_recent_news(company_name, max_results=max_results)
    except Exception as e:
        st.error(f"News search failed: {str(e)}")
        return []

def generate_with_ai(section, company_name, recent_news=None, provider=None):
    """Generate content using configured AI provider"""
    try:
        ai_provider = get_ai_provider(provider)
        return ai_provider.generate_content(section, company_name, recent_news)
    except Exception as e:
        return f"Error generating {section}: {str(e)}"

def create_word_document(company_name, content, recent_news):
    """Create Word document with new structure and return as bytes"""
    
    doc = Document()
    
    # Improve readability for the generated document
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(12)
    normal_style.paragraph_format.line_spacing = 1.15
    normal_style.paragraph_format.space_after = Pt(6)

    heading_1 = doc.styles['Heading 1']
    heading_1.font.name = 'Calibri'
    heading_1.font.size = Pt(16)
    heading_1.font.bold = True

    heading_2 = doc.styles['Heading 2']
    heading_2.font.name = 'Calibri'
    heading_2.font.size = Pt(13)
    heading_2.font.bold = True

    def set_paragraph_format(paragraph, font_size=12, bold=False, alignment=None):
        paragraph.style = 'Normal'
        paragraph.paragraph_format.space_after = Pt(6)
        paragraph.paragraph_format.line_spacing = 1.15
        if alignment is not None:
            paragraph.alignment = alignment
        for run in paragraph.runs:
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            run.bold = bold
        if not paragraph.runs:
            run = paragraph.add_run()
            run.font.name = 'Calibri'
            run.font.size = Pt(font_size)
            run.bold = bold

    # Title Page
    title = doc.add_heading(f"{company_name}", 0)
    title.alignment = 1  # Center
    title.style = 'Title'
    title.runs[0].font.name = 'Calibri'
    title.runs[0].font.size = Pt(24)
    title.runs[0].font.bold = True
    
    subtitle = doc.add_paragraph("Strategic Research Document")
    subtitle.alignment = 1
    subtitle.style = 'Subtitle'
    for run in subtitle.runs:
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
    
    generated_paragraph = doc.add_paragraph(f"Generated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
    set_paragraph_format(generated_paragraph, font_size=11)
    
    prepared_paragraph = doc.add_paragraph("Prepared by: AI Research Document Generator")
    set_paragraph_format(prepared_paragraph, font_size=11)
    
    divider = doc.add_paragraph("_" * 50)
    set_paragraph_format(divider, font_size=11)
    doc.add_page_break()
    
    # Table of Contents
    toc_heading = doc.add_heading("Table of Contents", level=1)
    set_paragraph_format(toc_heading, font_size=15, bold=True)
    sections = [
        "1. Executive Summary",
        "2. Company Overview",
        "3. Business Segments & Brands",
        "4. Sales & Distribution Model",
        "5. GCC Market Size, Growth & CAGR Analysis",
        "6. AI & Digital Transformation — Sales, Distribution & Retail",
        "7. Current Technology Stack",
        "8. Strategic Transformation & Future Planning",
        "9. Recent News Timeline (2024–2026)",
        "10. Key Insights: Growth Opportunities & Challenges",
        "11. Conclusion"
    ]
    for section in sections:
        p = doc.add_paragraph(section)
        set_paragraph_format(p, font_size=11)
    doc.add_page_break()
    
    # Section 1: Executive Summary
    doc.add_heading("1. Executive Summary", level=1)
    content_paragraph = doc.add_paragraph(content.get('executive_summary', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 2: Company Overview
    doc.add_heading("2. Company Overview", level=1)
    content_paragraph = doc.add_paragraph(content.get('company_overview', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 3: Business Segments & Brands
    doc.add_heading("3. Business Segments & Brands", level=1)
    content_paragraph = doc.add_paragraph(content.get('business_segments', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 4: Sales & Distribution Model
    doc.add_heading("4. Sales & Distribution Model", level=1)
    content_paragraph = doc.add_paragraph(content.get('sales_distribution', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 5: GCC Market Analysis
    doc.add_heading("5. GCC Market Size, Growth & CAGR Analysis", level=1)
    content_paragraph = doc.add_paragraph(content.get('gcc_market', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 6: AI & Digital Transformation
    doc.add_heading("6. AI & Digital Transformation — Sales, Distribution & Retail", level=1)
    content_paragraph = doc.add_paragraph(content.get('ai_transformation', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 7: Current Technology Stack
    doc.add_heading("7. Current Technology Stack", level=1)
    content_paragraph = doc.add_paragraph(content.get('tech_stack', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 8: Strategic Transformation
    doc.add_heading("8. Strategic Transformation & Future Planning", level=1)
    content_paragraph = doc.add_paragraph(content.get('strategic_transformation', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 9: Recent News Timeline
    doc.add_heading("9. Recent News Timeline (2024–2026)", level=1)
    content_paragraph = doc.add_paragraph(content.get('news_timeline', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    if recent_news:
        doc.add_heading("Recent News Sources:", level=2)
        for i, news in enumerate(recent_news, 1):
            p = doc.add_paragraph()
            p.add_run(f"{i}. ").bold = True
            p.add_run(news.get('title', 'No title'))
            p.add_run(f"\n   URL: {news.get('url', 'N/A')}\n")
            p.add_run(f"   {news.get('snippet', '')[:200]}...")
            set_paragraph_format(p, font_size=11)
    
    # Section 10: Key Insights
    doc.add_heading("10. Key Insights: Growth Opportunities & Challenges", level=1)
    content_paragraph = doc.add_paragraph(content.get('key_insights', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    # Section 11: Conclusion
    doc.add_heading("11. Conclusion", level=1)
    content_paragraph = doc.add_paragraph(content.get('conclusion', 'N/A'))
    set_paragraph_format(content_paragraph, font_size=12)
    
    footer_paragraph_1 = doc.add_paragraph("_" * 50)
    set_paragraph_format(footer_paragraph_1, font_size=11)
    footer_paragraph_2 = doc.add_paragraph("Generated by AI Research Document Generator")
    set_paragraph_format(footer_paragraph_2, font_size=11)
    footer_paragraph_3 = doc.add_paragraph("This document combines AI-generated insights with real-time web data.")
    set_paragraph_format(footer_paragraph_3, font_size=11)
    
    # Save to bytes
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    
    return buffer

# Generate button
if st.button("🚀 Generate Research Document", type="primary"):
    if not company_name:
        st.warning("⚠️ Please enter a company name!")
    else:
        with st.spinner(f"🔍 Researching {company_name}... This takes 2-3 minutes."):
            try:
                # Progress bar
                progress_bar = st.progress(0)
                status_text = st.empty()
                
                # Step 1: Search news
                status_text.text("Step 1/11: Searching recent news...")
                recent_news = search_recent_news(company_name, max_results=5)
                progress_bar.progress(8)
                
                # Generate all sections
                content = {}
                
                status_text.text("Step 2/11: Generating Executive Summary...")
                content['executive_summary'] = generate_with_ai('executive_summary', company_name, provider=selected_provider)
                progress_bar.progress(15)
                
                status_text.text("Step 3/11: Generating Company Overview...")
                content['company_overview'] = generate_with_ai('company_overview', company_name, provider=selected_provider)
                progress_bar.progress(23)
                
                status_text.text("Step 4/11: Generating Business Segments & Brands...")
                content['business_segments'] = generate_with_ai('business_segments', company_name, provider=selected_provider)
                progress_bar.progress(30)
                
                status_text.text("Step 5/11: Generating Sales & Distribution Model...")
                content['sales_distribution'] = generate_with_ai('sales_distribution', company_name, provider=selected_provider)
                progress_bar.progress(38)
                
                status_text.text("Step 6/11: Generating GCC Market Analysis...")
                content['gcc_market'] = generate_with_ai('gcc_market', company_name, provider=selected_provider)
                progress_bar.progress(46)
                
                status_text.text("Step 7/11: Generating AI & Digital Transformation...")
                content['ai_transformation'] = generate_with_ai('ai_transformation', company_name, provider=selected_provider)
                progress_bar.progress(54)
                
                status_text.text("Step 8/11: Generating Technology Stack...")
                content['tech_stack'] = generate_with_ai('tech_stack', company_name, provider=selected_provider)
                progress_bar.progress(62)
                
                status_text.text("Step 9/11: Generating Strategic Transformation...")
                content['strategic_transformation'] = generate_with_ai('strategic_transformation', company_name, provider=selected_provider)
                progress_bar.progress(69)
                
                status_text.text("Step 10/11: Generating News Timeline...")
                content['news_timeline'] = generate_with_ai('news_timeline', company_name, recent_news, provider=selected_provider)
                progress_bar.progress(77)
                
                status_text.text("Step 11/11: Generating Key Insights & Conclusion...")
                content['key_insights'] = generate_with_ai('key_insights', company_name, provider=selected_provider)
                content['conclusion'] = generate_with_ai('conclusion', company_name, provider=selected_provider)
                progress_bar.progress(100)
                
                status_text.text("✅ Creating document...")
                
                # Create document
                doc_bytes = create_word_document(company_name, content, recent_news)
                
                st.success(f"✅ Research document for *{company_name}* generated successfully!")
                
                # Download button
                timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
                filename = f"{company_name.replace(' ', '')}_Research{timestamp}.docx"
                
                st.download_button(
                    label="📥 Download Document (.docx)",
                    data=doc_bytes,
                    file_name=filename,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    type="primary"
                )
                
                # Show recent news
                if recent_news:
                    st.markdown("---")
                    st.subheader("📰 Recent News Found:")
                    for i, news in enumerate(recent_news, 1):
                        with st.expander(f"{i}. {news.get('title', 'No title')}"):
                            st.markdown(f"*URL:* {news.get('url', 'N/A')}")
                            st.markdown(f"{news.get('snippet', 'No snippet')}")
                
            except Exception as e:
                st.error(f"❌ Error: {str(e)}")
                st.error("Please check that your API key for the selected provider is configured in .env")

# Footer
st.markdown("---")
st.markdown("*Built with ❤️ using Streamlit, Claude/OpenAI/Gemini APIs, and DuckDuckGo Search*")