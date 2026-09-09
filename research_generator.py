import ollama
from duckduckgo_search import DDGS
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from datetime import datetime
import time

class ResearchDocumentGenerator:
    def _init_(self):
        self.company_name = ""
        self.content = {}
        self.recent_news = []
    
    def search_recent_news(self, company_name, max_results=5):
        """Search for recent news using DuckDuckGo"""
        print(f"🔍 Searching recent news for {company_name}...")
        
        try:
            # Search for recent news
            query = f"{company_name} news 2026"
            
            with DDGS() as ddgs:
                results = list(ddgs.text(query, max_results=max_results))
            
            news_items = []
            for result in results[:5]:
                news_items.append({
                    'title': result.get('title', 'No title'),
                    'url': result.get('href', 'No URL'),
                    'snippet': result.get('body', 'No snippet')
                })
            
            self.recent_news = news_items
            return news_items
        except Exception as e:
            print(f"⚠️ News search failed: {str(e)}")
            return []
    
    def generate_with_ai(self, section, company_name):
        """Generate content using Ollama"""
        
        prompts = {
            'overview': f"""Write a professional company overview for {company_name} (200-250 words).
            Include:
            - What the company does
            - Industry and sector
            - Headquarters location
            - Year founded
            - Key business areas
            - Company size (employees)
            
            Write in professional business language.""",
            
            'products': f"""List the key products and services offered by {company_name}.
            Format as bullet points with brief descriptions.
            Include 5-7 main products/services.
            For each product, mention:
            - Product name
            - What it does
            - Target customers
            
            Write in professional business language.""",
            
            'market': f"""Describe {company_name}'s market position and competitive landscape (200-250 words).
            Include:
            - Market share and ranking
            - Main competitors (name 3-5 competitors)
            - Target customer segments
            - Competitive advantages
            - Market trends affecting the company
            
            Write in professional business language.""",
            
            'sales': f"""Describe {company_name}'s sales position and strategy (200-250 words).
            Include:
            - Annual revenue (if known) or revenue range
            - Sales channels (direct, partners, online, etc.)
            - Sales strategy approach
            - Key sales markets/regions
            - Sales performance trends
            - Customer acquisition approach
            
            Write in professional business language.""",
            
            'challenges': f"""Analyze challenges in sales AI automation for {company_name} (250-300 words).
            Include:
            - Current AI adoption level in sales
            - Key challenges in implementing AI automation
            - Specific sales processes that need automation
            - Technology gaps or limitations
            - Data and integration challenges
            - Change management and adoption barriers
            - Recommendations for improvement
            
            Write in professional business language with actionable insights.""",
            
            'news_summary': f"""Based on these recent news items about {company_name}, write a brief summary (150 words):
            
            {self.recent_news}
            
            Summarize the key developments and trends. Write in professional business language."""
        }
        
        try:
            response = ollama.chat(
                model='llama3.2',
                messages=[{'role': 'user', 'content': prompts[section]}]
            )
            return response['message']['content']
        except Exception as e:
            return f"Error generating {section}: {str(e)}"
    
    def generate_research(self, company_name):
        """Generate complete research document"""
        self.company_name = company_name
        self.content = {}
        self.recent_news = []
        
        print(f"\n🔍 Researching {company_name}...")
        print("=" * 50)
        
        # Step 1: Search recent news (real-time)
        self.search_recent_news(company_name, max_results=5)
        
        # Step 2: Generate AI content (local, no limits)
        print("📝 Generating company overview...")
        self.content['overview'] = self.generate_with_ai('overview', company_name)
        
        print("📝 Generating products & services...")
        self.content['products'] = self.generate_with_ai('products', company_name)
        
        print("📝 Generating market position...")
        self.content['market'] = self.generate_with_ai('market', company_name)
        
        print("📝 Generating sales position...")
        self.content['sales'] = self.generate_with_ai('sales', company_name)
        
        print("📝 Generating sales AI automation challenges...")
        self.content['challenges'] = self.generate_with_ai('challenges', company_name)
        
        print("📝 Generating news summary...")
        self.content['news_summary'] = self.generate_with_ai('news_summary', company_name)
        
        print("=" * 50)
        print(f"✅ Content generation complete for {company_name}!")
        
        return self.content
    
    def create_word_document(self, output_filename=None):
        """Create professional Word document"""
        
        if not output_filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            output_filename = f"{self.company_name.replace(' ', '')}_Research{timestamp}.docx"
        
        print(f"\n📄 Creating Word document: {output_filename}")
        
        doc = Document()
        
        # Title Page
        title = doc.add_heading(f"{self.company_name}", 0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        subtitle = doc.add_paragraph("Research Document")
        subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        subtitle.style.font.size = Pt(16)
        subtitle.style.font.bold = True
        
        doc.add_paragraph(f"\nGenerated: {datetime.now().strftime('%B %d, %Y at %I:%M %p')}")
        doc.add_paragraph("Prepared by: AI Research Document Generator")
        doc.add_paragraph("_" * 50)
        
        # Table of Contents
        doc.add_heading("Table of Contents", level=1)
        doc.add_paragraph("1. Company Overview")
        doc.add_paragraph("2. Products & Services")
        doc.add_paragraph("3. Market Position")
        doc.add_paragraph("4. Sales Position")
        doc.add_paragraph("5. Challenges in Sales AI Automation")
        doc.add_paragraph("6. Recent News & Developments")
        doc.add_page_break()
        
        # Section 1: Company Overview
        doc.add_heading("1. Company Overview", level=1)
        doc.add_paragraph(self.content.get('overview', 'N/A'))
        
        # Section 2: Products & Services
        doc.add_heading("2. Products & Services", level=1)
        doc.add_paragraph(self.content.get('products', 'N/A'))
        
        # Section 3: Market Position
        doc.add_heading("3. Market Position", level=1)
        doc.add_paragraph(self.content.get('market', 'N/A'))
        
        # Section 4: Sales Position
        doc.add_heading("4. Sales Position", level=1)
        doc.add_paragraph(self.content.get('sales', 'N/A'))
        
        # Section 5: Challenges in Sales AI Automation
        doc.add_heading("5. Challenges in Sales AI Automation", level=1)
        doc.add_paragraph(self.content.get('challenges', 'N/A'))
        
        # Section 6: Recent News
        doc.add_heading("6. Recent News & Developments", level=1)
        
        # News summary
        doc.add_paragraph(self.content.get('news_summary', 'N/A'))
        
        # Recent news items
        if self.recent_news:
            doc.add_heading("Recent News Sources:", level=2)
            for i, news in enumerate(self.recent_news, 1):
                p = doc.add_paragraph()
                p.add_run(f"{i}. ").bold = True
                p.add_run(news.get('title', 'No title'))
                p.add_run(f"\n   URL: {news.get('url', 'N/A')}\n")
                p.add_run(f"   {news.get('snippet', '')[:200]}...")
        
        doc.add_paragraph("\n" + "_" * 50)
        doc.add_paragraph("Generated by AI Research Document Generator (Ollama + Web Search)")
        doc.add_paragraph("This document combines AI-generated insights with real-time web data.")
        
        doc.save(output_filename)
        print(f"✅ Document saved: {output_filename}")
        
        return output_filename
    
    def run(self, company_name):
        """Run complete research workflow"""
        self.generate_research(company_name)
        output_file = self.create_word_document()
        print(f"\n🎉 Research document for {self.company_name} completed!")
        print(f"📁 Output file: {output_file}")
        return output_file

if __name__ == "__main__":
    print("=" * 60)
    print("AI Research Document Generator")
    print("Powered by Ollama (Local AI) + Web Search")
    print("=" * 60)
    
    generator = ResearchDocumentGenerator()
    
    # Get company name
    company = input("\nEnter company name: ").strip()
    
    if not company:
        print("❌ Please enter a company name!")
        exit()
    
    # Generate research document
    output_file = generator.run(company)
    
    print("\n" + "=" * 60)
    print("✅ SUCCESS!")
    print(f"Document created: {output_file}")
    print("=" * 60)