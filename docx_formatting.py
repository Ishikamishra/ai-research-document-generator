"""Shared formatting helpers for generated research documents."""

import re
from datetime import datetime

from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SECTION_TITLES = [
    "Company Overview",
    "Products & Services",
    "Market Position",
    "Sales Position",
    "Challenges in Sales AI Automation",
    "Recent News & Developments",
]

CONTENT_KEYS = [
    "overview",
    "products",
    "market",
    "sales",
    "challenges",
    "news_summary",
]


def _set_cell_shading(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _set_repeat_table_header(row):
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def configure_document(doc):
    """Apply consistent page, font, heading, header, and footer styles."""
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for name, size, color in (("Title", 28, "17365D"), ("Heading 1", 17, "17365D"), ("Heading 2", 12, "2F75B5")):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = "AI RESEARCH DOCUMENT GENERATOR"
    header.style = styles["Caption"]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    footer = section.footer.paragraphs[0]
    _set_page_number(footer)
    footer.style = styles["Caption"]


def add_title_page(doc, company_name):
    title = doc.add_heading(company_name, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(90)

    subtitle = doc.add_paragraph("Company Research Brief")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.runs[0].font.size = Pt(16)
    subtitle.runs[0].font.color.rgb = RGBColor(47, 117, 181)

    metadata = doc.add_paragraph()
    metadata.alignment = WD_ALIGN_PARAGRAPH.CENTER
    metadata.paragraph_format.space_before = Pt(20)
    metadata.add_run("Generated ").bold = True
    metadata.add_run(datetime.now().strftime("%B %d, %Y"))

    doc.add_paragraph(
        "This report combines AI-generated analysis with current web research.",
        style="Subtitle",
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def add_table_of_contents(doc):
    doc.add_heading("Contents", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light Shading Accent 1"
    table.autofit = False
    table.columns[0].width = Inches(0.65)
    table.columns[1].width = Inches(5.9)
    header = table.rows[0]
    header.cells[0].text = "No."
    header.cells[1].text = "Section"
    _set_repeat_table_header(header)
    for cell in header.cells:
        _set_cell_shading(cell, "17365D")
        for run in cell.paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)

    for index, title in enumerate(SECTION_TITLES, 1):
        cells = table.add_row().cells
        cells[0].text = str(index)
        cells[1].text = title
    doc.add_page_break()


def _add_markdown_runs(paragraph, text):
    """Convert the common Markdown emphasis markers used by the models."""
    pattern = re.compile(r"(\*\*|__)(.+?)\1")
    position = 0
    for match in pattern.finditer(text):
        if match.start() > position:
            paragraph.add_run(text[position:match.start()])
        run = paragraph.add_run(match.group(2))
        run.bold = True
        position = match.end()

    if position < len(text):
        paragraph.add_run(text[position:])


def _is_nested_detail(text):
    return bool(
        re.match(
            r"^\*\*(what it does|target customers|key features|benefits)\s*:",
            text,
            flags=re.IGNORECASE,
        )
    )


def _is_bullet_heading(text):
    return bool(re.match(r"^\*\*[^*]+\*\*:?$", text))


def add_ai_content(doc, content):
    """Render model output as readable paragraphs and real Word bullets."""
    text = str(content or "N/A").replace("\r\n", "\n")
    for raw_line in text.split("\n"):
        line = raw_line.strip()
        if not line:
            continue

        bullet_match = re.match(r"^(?:[-*•]|\d+[.)])\s+(.*)$", line)
        if bullet_match:
            bullet_text = bullet_match.group(1).strip()
            style = "List Bullet 2" if _is_nested_detail(bullet_text) else "List Bullet"
            paragraph = doc.add_paragraph(style=style)
            if style == "List Bullet" and _is_bullet_heading(bullet_text):
                paragraph.paragraph_format.keep_with_next = True
            _add_markdown_runs(paragraph, bullet_text)
        else:
            paragraph = doc.add_paragraph()
            _add_markdown_runs(paragraph, line)


def add_news_sources(doc, recent_news):
    if not recent_news:
        return

    doc.add_heading("Sources", level=2)
    for index, news in enumerate(recent_news, 1):
        paragraph = doc.add_paragraph(style="List Number")
        paragraph.add_run(news.get("title", "Untitled source")).bold = True
        url = news.get("url", "")
        if url:
            paragraph.add_run(f"\n{url}")
        snippet = news.get("snippet", "").strip()
        if snippet:
            paragraph.add_run(f"\n{snippet[:300]}")


def add_section(doc, number, title, content):
    doc.add_heading(f"{number}. {title}", level=1)
    add_ai_content(doc, content)
